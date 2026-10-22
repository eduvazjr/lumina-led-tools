import openpyxl
from datetime import datetime, date
from decimal import Decimal
import re
import os
import sys

# Adicionar o caminho do Django ao sys.path
sys.path.insert(0, r'C:\Users\Usuario\djangoSIGE')
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'djangosige.configs.settings')

import django
django.setup()

from djangosige.apps.gerencial.models import PagamentoFornecedor
from djangosige.apps.cadastro.models import Fornecedor
from django.contrib.auth import get_user_model

User = get_user_model()

# Função para limpar valor monetário
def limpar_valor(valor_str):
    """Converte string monetária para Decimal"""
    if not valor_str:
        return Decimal('0')
    
    if isinstance(valor_str, (int, float)):
        return Decimal(str(valor_str))
    
    # Remover R$, espaços, pontos (milhares) e substituir vírgula por ponto
    valor_limpo = str(valor_str).replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
    
    try:
        return Decimal(valor_limpo)
    except:
        return Decimal('0')

# Função para converter data
def converter_data(data_val):
    """Converte valor de data para datetime.date"""
    if isinstance(data_val, datetime):
        return data_val.date()
    elif isinstance(data_val, date):
        return data_val
    elif isinstance(data_val, str):
        try:
            return datetime.strptime(data_val, '%d/%m/%Y').date()
        except:
            return date.today()
    return date.today()

# Pegar usuário admin
try:
    user = User.objects.get(username='admin')
except:
    user = User.objects.first()

# Abrir planilha
wb = openpyxl.load_workbook(r'C:\Users\Usuario\Desktop\Controle Fornecedores (1).xlsx')

# Mapear abas para fornecedores
fornecedores_map = {
    'Amarante': 'Amarante',
    'Nakasaki': 'Nakasaki',
    'Lucas': 'Lucas',
    'Terra 2025': 'Terra',
    'Terra 2026': 'Terra',
    'Max': 'Max',
    'Yé': 'Yé'
}

print("="*70)
print("🚀 IMPORTAÇÃO DE PAGAMENTOS PARA O CONTROLE GERENCIAL")
print("="*70)

total_importados = 0
total_erros = 0
fornecedores_criados = []

for aba_nome, fornecedor_nome in fornecedores_map.items():
    if aba_nome not in wb.sheetnames:
        continue
    
    ws = wb[aba_nome]
    
    print(f"\n📄 Processando aba: {aba_nome}")
    print("-" * 70)
    
    # Buscar ou criar fornecedor
    fornecedor, criado = Fornecedor.objects.get_or_create(
        nome_razao_social=fornecedor_nome,
        defaults={
            'tipo_pessoa': 'PJ',
        }
    )
    
    if criado:
        fornecedores_criados.append(fornecedor_nome)
        print(f"   ✅ Fornecedor '{fornecedor_nome}' criado")
    else:
        print(f"   ℹ️  Fornecedor '{fornecedor_nome}' já existe")
    
    # Processar linhas
    importados_aba = 0
    erros_aba = 0
    
    # Encontrar linha de cabeçalho
    header_row = None
    for row_idx in range(1, 20):
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row_idx, col).value
            if val and 'data' in str(val).lower():
                header_row = row_idx
                break
        if header_row:
            break
    
    if not header_row:
        print(f"   ⚠️  Cabeçalho não encontrado, pulando...")
        continue
    
    # Identificar colunas
    col_data_debito = None
    col_desc_debito = None
    col_valor_debito = None
    
    for col in range(1, ws.max_column + 1):
        header = ws.cell(header_row, col).value
        if not header:
            continue
        
        header_str = str(header).lower().strip()
        
        # Primeira coluna "Data" é débito
        if 'data' in header_str and col_data_debito is None:
            col_data_debito = col
        # Primeira coluna "Descrição" é débito  
        elif 'descri' in header_str and col_desc_debito is None:
            col_desc_debito = col
        # Coluna "Valor" próxima da descrição de débito
        elif 'valor' in header_str and col_desc_debito and col_valor_debito is None:
            col_valor_debito = col
    
    print(f"   📊 Colunas identificadas: Data={col_data_debito}, Desc={col_desc_debito}, Valor={col_valor_debito}")
    
    if not all([col_data_debito, col_desc_debito, col_valor_debito]):
        print(f"   ⚠️  Colunas não identificadas, pulando...")
        continue
    
    # Processar linhas de dados
    for row_idx in range(header_row + 1, ws.max_row + 1):
        try:
            data_val = ws.cell(row_idx, col_data_debito).value
            desc_val = ws.cell(row_idx, col_desc_debito).value
            valor_val = ws.cell(row_idx, col_valor_debito).value
            
            # Pular linhas vazias
            if not valor_val or valor_val == 0:
                continue
            
            # Converter valores
            data_emissao = converter_data(data_val)
            descricao = str(desc_val) if desc_val else f'Compra {aba_nome}'
            valor = limpar_valor(valor_val)
            
            if valor <= 0:
                continue
            
            # Data de vencimento: 30 dias após emissão
            from datetime import timedelta
            data_vencimento = data_emissao + timedelta(days=30)
            
            # Criar pagamento
            PagamentoFornecedor.objects.create(
                fornecedor=fornecedor,
                descricao=descricao[:200],
                valor=valor,
                data_emissao=data_emissao,
                data_vencimento=data_vencimento,
                status='pendente',
                observacoes=f'Importado da aba: {aba_nome}',
                usuario=user
            )
            
            importados_aba += 1
            total_importados += 1
            
        except Exception as e:
            erros_aba += 1
            total_erros += 1
            #print(f"      ❌ Erro linha {row_idx}: {str(e)}")
    
    print(f"   ✅ {importados_aba} pagamentos importados | ❌ {erros_aba} erros")

print("\n" + "="*70)
print("🎉 IMPORTAÇÃO CONCLUÍDA!")
print("="*70)
print(f"\n✅ Total de pagamentos importados: {total_importados}")
print(f"❌ Total de erros: {total_erros}")

if fornecedores_criados:
    print(f"\n👤 Fornecedores criados ({len(fornecedores_criados)}):")
    for nome in fornecedores_criados:
        print(f"   • {nome}")

print("\n🌐 Acesse: http://localhost:8000/gerencial/")
print("="*70)
