import openpyxl
from datetime import datetime
import re

# Abrir a planilha
wb = openpyxl.load_workbook(r'C:\Users\Usuario\Desktop\Controle Fornecedores (1).xlsx')

print("="*70)
print("📊 ANÁLISE DETALHADA DA PLANILHA")
print("="*70)

# Analisar cada aba
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n📄 ABA: {sheet_name}")
    print("-" * 70)
    
    # Procurar linha com cabeçalhos válidos
    header_row = None
    for row_idx in range(1, min(20, ws.max_row + 1)):
        row_values = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
        row_text = ' '.join([str(v) for v in row_values if v])
        
        # Procurar por palavras-chave de cabeçalhos
        keywords = ['pedido', 'produto', 'valor', 'vencimento', 'parcela', 'quantidade']
        if any(keyword in row_text.lower() for keyword in keywords):
            header_row = row_idx
            print(f"   ✅ Cabeçalhos encontrados na linha {row_idx}")
            
            # Mostrar cabeçalhos
            headers = []
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row_idx, col).value
                if val and str(val).strip():
                    headers.append((col, str(val).strip()))
            
            print(f"   📌 Colunas identificadas:")
            for col_idx, header in headers:
                print(f"      Col {col_idx}: {header}")
            
            # Contar linhas com dados
            data_rows = 0
            for row_idx in range(header_row + 1, ws.max_row + 1):
                # Verificar se tem algum dado relevante
                has_data = False
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row_idx, col).value
                    if val and str(val).strip() and str(val).strip() != 'None':
                        has_data = True
                        break
                if has_data:
                    data_rows += 1
            
            print(f"   📊 Linhas com dados: {data_rows}")
            
            # Mostrar 3 primeiras linhas de exemplo
            print(f"   📝 Primeiras 3 linhas de exemplo:")
            for row_idx in range(header_row + 1, min(header_row + 4, ws.max_row + 1)):
                print(f"\n      Linha {row_idx}:")
                for col_idx, header in headers[:8]:  # Mostrar primeiras 8 colunas
                    val = ws.cell(row_idx, col_idx).value
                    if isinstance(val, datetime):
                        val = val.strftime('%d/%m/%Y')
                    if val:
                        print(f"         {header}: {val}")
            
            break
    
    if not header_row:
        print(f"   ⚠️  Cabeçalhos não identificados automaticamente")
        print(f"   Linhas totais: {ws.max_row}")

print("\n" + "="*70)
print("✅ ANÁLISE CONCLUÍDA")
print("="*70)
