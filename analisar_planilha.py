import openpyxl
from datetime import datetime
import sys

# Abrir a planilha
print("Abrindo planilha...")
wb = openpyxl.load_workbook(r'C:\Users\Usuario\Desktop\Controle Fornecedores (1).xlsx')

# Listar todas as abas
print("\n" + "="*60)
print("📊 ABAS DISPONÍVEIS NA PLANILHA:")
print("="*60)
for sheet_name in wb.sheetnames:
    print(f"   📄 {sheet_name}")

print("\n" + "="*60)
print("📋 ESTRUTURA DA PRIMEIRA ABA:")
print("="*60)

# Ler a primeira aba
ws = wb.active
print(f"\nAba ativa: {ws.title}")
print(f"Total de linhas: {ws.max_row}")
print(f"Total de colunas: {ws.max_column}")

# Mostrar cabeçalhos (primeira linha)
print("\n📌 CABEÇALHOS (Linha 1):")
headers = []
for col in range(1, ws.max_column + 1):
    cell_value = ws.cell(1, col).value
    headers.append(cell_value)
    print(f"   Coluna {col}: {cell_value}")

# Mostrar primeiras 5 linhas de dados
print("\n📝 PRIMEIRAS 5 LINHAS DE DADOS:")
print("="*60)
for row_idx in range(2, min(7, ws.max_row + 1)):
    print(f"\nLinha {row_idx}:")
    for col_idx, header in enumerate(headers, start=1):
        cell_value = ws.cell(row_idx, col_idx).value
        if isinstance(cell_value, datetime):
            cell_value = cell_value.strftime('%d/%m/%Y')
        print(f"   {header}: {cell_value}")

print("\n" + "="*60)
print(f"📊 TOTAL DE REGISTROS: {ws.max_row - 1}")
print("="*60)
